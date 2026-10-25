"""Тесты ReportModule: отчёты, права доступа, экспорт."""
import io
from datetime import date, timedelta

from openpyxl import load_workbook

from tests.base import ApiTestCase


class ReportTests(ApiTestCase):
    """Отчёты по исполнителю, проекту, периоду и экспорт в XLSX/PDF."""

    def setUp(self):
        super().setUp()
        self.manager, self.mh = self.make_user("manager")
        self.exe, self.eh = self.make_user("executor")
        self.exe2, self.eh2 = self.make_user("executor")
        self.p1 = self.make_project(self.mh, "Проект 1")
        self.p2 = self.make_project(self.mh, "Проект 2")
        self.t1 = self.make_task(self.mh, self.p1, title="Задача А", assignee_id=self.exe["id"])
        self.t2 = self.make_task(self.mh, self.p1, title="Задача Б", assignee_id=self.exe["id"])
        self.t3 = self.make_task(self.mh, self.p2, title="Задача В", assignee_id=self.exe2["id"])
        self.set_status(self.eh, self.t1["id"], "in_progress")

    def test_report_by_assignee(self):
        resp = self.client.get(f"/api/reports/assignee?assignee_id={self.exe['id']}",
                               headers=self.mh)
        data = resp.get_json()
        self.assertEqual(data["summary"]["total"], 2)
        self.assertEqual(data["summary"]["by_status"]["in_progress"], 1)
        self.assertEqual(data["summary"]["by_status"]["new"], 1)

    def test_report_by_project(self):
        data = self.client.get(f"/api/reports/project?project_id={self.p2}",
                               headers=self.mh).get_json()
        self.assertEqual([r["title"] for r in data["rows"]], ["Задача В"])
        self.assertEqual(data["rows"][0]["project_name"], "Проект 2")

    def test_report_by_period_inclusive_bounds(self):
        today = date.today().isoformat()
        data = self.client.get(f"/api/reports/period?date_from={today}&date_to={today}",
                               headers=self.mh).get_json()
        self.assertEqual(data["summary"]["total"], 3)
        future = (date.today() + timedelta(days=5)).isoformat()
        empty = self.client.get(f"/api/reports/period?date_from={future}&date_to={future}",
                                headers=self.mh).get_json()
        self.assertEqual(empty["summary"]["total"], 0)
        self.assertEqual(empty["rows"], [])

    def test_report_param_validation(self):
        bad_urls = [
            "/api/reports/assignee",
            "/api/reports/project",
            "/api/reports/period?date_from=2026-01-01",
            "/api/reports/period?date_from=2026-02-01&date_to=2026-01-01",
            "/api/reports/period?date_from=bad&date_to=2026-01-01",
            "/api/reports/assignee?assignee_id=1&format=docx",
        ]
        for url in bad_urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url, headers=self.mh).status_code, 400)

    def test_unknown_report_type_and_unknown_project(self):
        self.assertEqual(self.client.get("/api/reports/other", headers=self.mh).status_code, 404)
        self.assertEqual(self.client.get("/api/reports/project?project_id=999",
                                         headers=self.mh).status_code, 404)

    def test_executor_gets_only_own_report(self):
        own = self.client.get("/api/reports/assignee", headers=self.eh).get_json()
        self.assertEqual(own["summary"]["total"], 2)
        foreign = self.client.get(f"/api/reports/assignee?assignee_id={self.exe2['id']}",
                                  headers=self.eh)
        self.assertEqual(foreign.status_code, 403)
        self.assertEqual(self.client.get(f"/api/reports/project?project_id={self.p1}",
                                         headers=self.eh).status_code, 403)

    def test_history_is_manager_only_and_records_reports(self):
        self.client.get(f"/api/reports/project?project_id={self.p1}", headers=self.mh)
        history = self.client.get("/api/reports", headers=self.mh).get_json()
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["rows_count"], 2)
        self.assertEqual(self.client.get("/api/reports", headers=self.eh).status_code, 403)

    def test_overdue_counted(self):
        with self.app.app_context():
            from app.db import get_db
            db = get_db()
            db.execute("UPDATE tasks SET due_date = '2020-01-01' WHERE id = ?",
                       (self.t2["id"],))
            db.commit()
        data = self.client.get(f"/api/reports/project?project_id={self.p1}",
                               headers=self.mh).get_json()
        self.assertEqual(data["summary"]["overdue"], 1)

    def test_export_xlsx(self):
        resp = self.client.get(f"/api/reports/project?project_id={self.p1}&format=xlsx",
                               headers=self.mh)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp.mimetype)
        ws = load_workbook(io.BytesIO(resp.data))["Задачи"]
        self.assertEqual(ws.max_row, 3)          # заголовок + 2 задачи
        self.assertEqual(ws["C2"].value, "Задача А")

    def test_export_pdf(self):
        resp = self.client.get(f"/api/reports/project?project_id={self.p1}&format=pdf",
                               headers=self.mh)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.mimetype, "application/pdf")
        self.assertTrue(resp.data.startswith(b"%PDF"))

    def test_export_empty_report_pdf_and_xlsx(self):
        future = (date.today() + timedelta(days=9)).isoformat()
        for fmt in ("pdf", "xlsx"):
            resp = self.client.get(
                f"/api/reports/period?date_from={future}&date_to={future}&format={fmt}",
                headers=self.mh)
            self.assertEqual(resp.status_code, 200)
