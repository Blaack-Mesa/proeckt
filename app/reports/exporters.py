"""Экспорт отчёта в Excel (openpyxl) и PDF (reportlab)."""
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COLUMNS = [("id", "№"), ("project_name", "Проект"), ("title", "Задача"),
           ("assignee_name", "Исполнитель"), ("status_label", "Статус"),
           ("due_date", "Срок")]

# Шрифт с кириллицей; при отсутствии используется встроенный Helvetica.
_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def _register_font() -> str:
    """Регистрирует TTF-шрифт и возвращает его имя."""
    for path in _FONT_CANDIDATES:
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont("PFFont", path))
            return "PFFont"
    return "Helvetica"


def _cell(row: dict, key: str) -> str:
    """Значение ячейки как строка (пустое для None)."""
    value = row.get(key)
    return "" if value is None else str(value)


def to_xlsx(report: dict) -> bytes:
    """Excel-файл: лист «Задачи» (таблица) и лист «Сводка»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Задачи"
    ws.append([title for _, title in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in report["rows"]:
        ws.append([_cell(row, key) for key, _ in COLUMNS])
    for idx, width in enumerate((6, 24, 40, 24, 16, 12), start=1):
        ws.column_dimensions[chr(64 + idx)].width = width

    summary = wb.create_sheet("Сводка")
    summary.append([report["title"]])
    summary["A1"].font = Font(bold=True)
    summary.append(["Сформирован", report["generated_at"]])
    summary.append(["Всего задач", report["summary"]["total"]])
    summary.append(["Просрочено", report["summary"]["overdue"]])
    for status, count in report["summary"]["by_status"].items():
        summary.append([status, count])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(report: dict) -> bytes:
    """PDF-файл: заголовок, сводка и таблица задач."""
    font = _register_font()
    styles = getSampleStyleSheet()
    for name in ("Title", "Normal"):
        styles[name].fontName = font
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=report["title"])
    summary = report["summary"]
    story = [
        Paragraph(report["title"], styles["Title"]),
        Paragraph(f"Сформирован: {report['generated_at']}. Всего задач: {summary['total']}, "
                  f"просрочено: {summary['overdue']}", styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [[title for _, title in COLUMNS]]
    data += [[_cell(row, key) for key, _ in COLUMNS] for row in report["rows"]]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(table)
    doc.build(story)
    return buffer.getvalue()
